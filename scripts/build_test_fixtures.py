"""
임포트 엔진 검증용 '지저분한' 엑셀 픽스처 생성기.

    py scripts/build_test_fixtures.py

실제 축구교실 파일은 A1부터 깔끔하게 시작하지 않는다. 제목이 붙고, 헤더가
4행부터고, 반별 블록이 세로로 쌓이고, 헤더가 중간에 반복되고, 맨 아래
합계와 '※ 문의'가 달린다. 우리가 만든 양식만 읽히는 엔진은 영업에서 죽으므로,
일부러 서로 다른 방식으로 어긋난 6종을 만들어 무설정 통과율을 잰다.

산출: src/lib/import/__fixtures__/*.xlsx
"""

from __future__ import annotations

import os
from datetime import date

from openpyxl import Workbook

OUT_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "src", "lib", "import", "__fixtures__",
)

NAMES = [
    "김서준", "이도윤", "박하준", "최지호", "정예준", "강시우", "조주원", "윤지훈",
    "장건우", "임서연", "한유진", "오민재", "신준호", "권태양", "황라온", "안수빈",
]
PARENTS = [
    "김철수", "이영희", "박민수", "최수진", "정대현", "강미경", "조성훈", "윤경석",
    "장하영", "임동주", "한지민", "오세영", "신부영", "권나래", "황보경", "안정호",
]

assert all(all("가" <= ch <= "힣" for ch in n) for n in NAMES + PARENTS), "이름 풀에 한글 아닌 문자"


def phone(i: int) -> str:
    return f"010-{2000 + i * 7:04d}-{1000 + i * 13:04d}"


def student(i: int, cls: str, fee: int = 150000):
    """(이름, 반, 연락처, 수강료, 등록일, 생년월일, 보호자)"""
    return (
        NAMES[i % len(NAMES)],
        cls,
        phone(i),
        fee,
        date(2024 + (i % 2), 1 + (i * 3) % 12, 1 + (i * 5) % 28),
        date(2015 + (i % 5), 1 + (i * 7) % 12, 1 + (i * 11) % 28),
        PARENTS[i % len(PARENTS)],
    )


# ---------------------------------------------------------------------------
# 1. 깔끔한 표 — 기준선
# ---------------------------------------------------------------------------
def fixture_clean(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "원생명단"
    ws.append(["이름", "반", "학부모연락처", "수강료", "등록일"])
    for i in range(12):
        n, c, p, f, e, _b, _pa = student(i, "U9 화목반" if i % 2 else "U11 월수금반")
        ws.append([n, c, p, f, e])
    wb.save(path)


# ---------------------------------------------------------------------------
# 2. 제목 행 + 헤더 4행부터 + 합계 + 꼬리말, 헤더명이 전부 다름
# ---------------------------------------------------------------------------
def fixture_titled(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "2026 명단"
    ws.append(["성동FC 유소년 원생 명단"])
    ws.append(["2026학년도 1학기 기준"])
    ws.append([])
    ws.append(["성명", "소속", "연락처", "월회비", "가입일", "생년월일", "학부모명"])
    for i in range(14):
        ws.append(list(student(i, "U9 화목반" if i % 3 else "U13 월수반", 160000)))
    ws.append([])
    ws.append(["합계", "", "", 2240000])
    ws.append(["※ 문의: 010-1111-2222"])
    wb.save(path)


# ---------------------------------------------------------------------------
# 3. 반별 블록이 세로로 쌓임 + 헤더 반복 (반 이름은 구분 행에만 있음)
# ---------------------------------------------------------------------------
def fixture_sections(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "반별명단"
    header = ["이름", "보호자연락처", "수강료", "등록일"]
    idx = 0
    for cls in ["U7 토요반", "U9 화목반", "U11 월수금반"]:
        ws.append([cls])
        ws.append(header)
        for _ in range(5):
            n, _c, p, f, e, _b, _pa = student(idx, cls, 140000)
            ws.append([n, p, f, e])
            idx += 1
        ws.append([])
    ws.append(["총계", "", 2100000])
    wb.save(path)


# ---------------------------------------------------------------------------
# 4. 헤더가 아예 없음 — 내용만으로 판별해야 한다
# ---------------------------------------------------------------------------
def fixture_headerless(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for i in range(10):
        n, c, p, f, e, _b, _pa = student(i, "U9 화목반" if i % 2 else "U13 월수반")
        ws.append([n, c, p, f, e])
    wb.save(path)


# ---------------------------------------------------------------------------
# 5. 납부 그리드 + 계산 컬럼(무시 대상) + 왼쪽 빈 열
# ---------------------------------------------------------------------------
def fixture_payments(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "수강생"
    ws.append([None, "이름", "반", "연락처", "월회비",
               "2026-06", "2026-07", "2026-08", "상태", "이탈위험도"])
    marks = [["O", "O", "O"], ["O", "O", None], [None, "X", "O"], ["O", "O", "150000"]]
    for i in range(12):
        n, c, p, f, _e, _b, _pa = student(i, "U9 화목반" if i % 2 else "U11 월수금반")
        ws.append([None, n, c, p, f, *marks[i % len(marks)], "정상", 12])
    wb.save(path)


# ---------------------------------------------------------------------------
# 6. 시트 하나가 반 하나 — 반 이름이 시트명에만 있음
# ---------------------------------------------------------------------------
def fixture_multisheet(path):
    wb = Workbook()
    wb.remove(wb.active)
    idx = 0
    for cls in ["U9 화목반", "U11 월수금반", "U13 월수반"]:
        ws = wb.create_sheet(cls)
        ws.append(["이름", "연락처", "수강료"])
        for _ in range(6):
            n, _c, p, f, _e, _b, _pa = student(idx, cls)
            ws.append([n, p, f])
            idx += 1
    wb.save(path)


FIXTURES = [
    ("01-clean.xlsx", fixture_clean, 12),
    ("02-titled.xlsx", fixture_titled, 14),
    ("03-sections.xlsx", fixture_sections, 15),
    ("04-headerless.xlsx", fixture_headerless, 10),
    ("05-payments.xlsx", fixture_payments, 12),
    ("06-multisheet.xlsx", fixture_multisheet, 18),
]


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    for filename, build, expected in FIXTURES:
        path = os.path.join(OUT_DIR, filename)
        build(path)
        print(f"written: {filename}  (expected students: {expected})")


if __name__ == "__main__":
    main()
